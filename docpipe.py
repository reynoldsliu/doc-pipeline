#!/usr/bin/env python3
"""docpipe — 會議逐字稿 → 標準化文件的 LLM pipeline。

設計原則：
1. Provider 以環境變數抽換（LLM_PROVIDER=claude|gemini|fake），零程式碼切換。
2. Prompt 模板是資料不是程式：templates/*.md 可自行增修，不動本檔。
3. fake provider 讓測試與 CI 完全離線，也讓新使用者零金鑰試跑。
"""
import argparse
import datetime
import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
TEMPLATE_DIR = ROOT / "templates"

SYSTEM_PROMPT = (
    "你是嚴謹的技術文件產生器。只輸出模板要求的 Markdown 內容，"
    "不要加任何前言、結語或程式碼圍欄。逐字稿沒提到的資訊寫「未提及」，不要編造。"
)


# ---------- templates ----------

def list_templates() -> list[str]:
    return sorted(p.stem for p in TEMPLATE_DIR.glob("*.md"))


def render(template: str, transcript: str) -> str:
    path = TEMPLATE_DIR / f"{template}.md"
    if not path.exists():
        raise SystemExit(f"找不到模板 '{template}'，可用：{', '.join(list_templates())}")
    today = datetime.date.today().isoformat()
    return path.read_text(encoding="utf-8").replace("{transcript}", transcript).replace("{today}", today)


# ---------- providers ----------

class FakeProvider:
    """離線 provider：測試、CI、無金鑰試跑用。輸出可預期、不連網。"""
    name = "fake"

    def __init__(self, model: str | None = None):
        self.model = model or "fake-1"

    def generate(self, prompt: str) -> str:
        head = prompt.strip().splitlines()[0] if prompt.strip() else ""
        return (
            "# 產出（fake provider，未呼叫任何 LLM）\n\n"
            f"- prompt 首行：{head[:80]}\n"
            f"- prompt 長度：{len(prompt)} 字元\n"
        )


class ClaudeProvider:
    name = "claude"

    def __init__(self, model: str | None = None):
        try:
            import anthropic
        except ImportError:
            raise SystemExit("claude provider 需要 anthropic 套件。\n"
                             "  python3 -m venv .venv && .venv/bin/pip install anthropic\n"
                             "  然後用 .venv/bin/python 執行本工具\n"
                             "（Homebrew/系統 Python 會拒絕直接 pip install，屬正常保護機制）")
        if not (os.environ.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_AUTH_TOKEN")):
            raise SystemExit("缺少 ANTHROPIC_API_KEY（或先 `ant auth login`）")
        self.client = anthropic.Anthropic()
        self.model = model or os.environ.get("DOCPIPE_CLAUDE_MODEL", "claude-opus-5")

    def generate(self, prompt: str) -> str:
        # 文件產出可能很長：走 streaming 避免 HTTP timeout
        with self.client.messages.stream(
            model=self.model,
            max_tokens=64000,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": prompt}],
        ) as stream:
            message = stream.get_final_message()
        return "".join(b.text for b in message.content if b.type == "text")


class GeminiProvider:
    name = "gemini"

    def __init__(self, model: str | None = None):
        try:
            from google import genai
        except ImportError:
            raise SystemExit("gemini provider 需要 google-genai 套件。\n"
                             "  python3 -m venv .venv && .venv/bin/pip install google-genai\n"
                             "  然後用 .venv/bin/python 執行本工具")
        if not os.environ.get("GEMINI_API_KEY"):
            raise SystemExit("缺少 GEMINI_API_KEY")
        self.client = genai.Client()
        self.model = model or os.environ.get("DOCPIPE_GEMINI_MODEL", "gemini-2.5-pro")

    def generate(self, prompt: str) -> str:
        resp = self.client.models.generate_content(
            model=self.model, contents=f"{SYSTEM_PROMPT}\n\n{prompt}"
        )
        return resp.text or ""


PROVIDERS = {"claude": ClaudeProvider, "gemini": GeminiProvider, "fake": FakeProvider}


def get_provider(model: str | None = None):
    name = os.environ.get("LLM_PROVIDER", "claude").lower()
    if name not in PROVIDERS:
        raise SystemExit(f"未知的 LLM_PROVIDER '{name}'，可用：{', '.join(PROVIDERS)}")
    return PROVIDERS[name](model)


# ---------- output ----------

def md_to_docx(md_text: str, out_path: Path) -> None:
    """極簡 md → docx：標題／清單／段落。要完整格式請用 pandoc。"""
    try:
        import docx
    except ImportError:
        raise SystemExit("--docx 需要 python-docx 套件：.venv/bin/pip install python-docx")
    d = docx.Document()
    for line in md_text.splitlines():
        s = line.strip()
        if not s or s.startswith("<!--"):
            continue
        if s.startswith("#"):
            level = min(len(s) - len(s.lstrip("#")), 4)
            d.add_heading(s.lstrip("#").strip(), level=level)
        elif s.startswith(("- ", "* ")):
            d.add_paragraph(s[2:], style="List Bullet")
        else:
            d.add_paragraph(s)
    d.save(out_path)


def run(files: list[str], template: str, outdir: str, docx: bool, model: str | None) -> list[Path]:
    provider = get_provider(model)
    out = Path(outdir)
    out.mkdir(parents=True, exist_ok=True)
    written = []
    for f in files:
        src = Path(f)
        transcript = src.read_text(encoding="utf-8")
        body = provider.generate(render(template, transcript))
        stamp = datetime.datetime.now().isoformat(timespec="seconds")
        footer = f"\n\n<!-- docpipe: provider={provider.name} model={provider.model} template={template} source={src.name} generated={stamp} -->\n"
        dest = out / f"{src.stem}.{template}.md"
        dest.write_text(body.rstrip() + footer, encoding="utf-8")
        written.append(dest)
        if docx:
            md_to_docx(body, dest.with_suffix(".docx"))
            written.append(dest.with_suffix(".docx"))
    return written


# ---------- cli ----------

def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="docpipe", description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list", help="列出可用模板")
    rp = sub.add_parser("run", help="產出文件")
    rp.add_argument("files", nargs="+", help="逐字稿檔案（txt/md）")
    rp.add_argument("-t", "--template", default="meeting-minutes")
    rp.add_argument("-o", "--out", default="out")
    rp.add_argument("--model", help="覆寫 provider 預設模型")
    rp.add_argument("--docx", action="store_true", help="同時輸出 .docx（需 python-docx）")
    args = ap.parse_args(argv)

    if args.cmd == "list":
        print("\n".join(list_templates()))
        return 0
    for p in run(args.files, args.template, args.out, args.docx, args.model):
        print(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
